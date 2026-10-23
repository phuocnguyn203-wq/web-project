import os
import tempfile

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp()}/test.db"
os.environ["UPLOAD_DIR"] = tempfile.mkdtemp()
os.environ["SYNC_JOBS"] = "1"

import io  # noqa: E402
import zipfile  # noqa: E402
from pathlib import Path  # noqa: E402

import pytest  # noqa: E402
from alembic import command  # noqa: E402
from alembic.config import Config  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pptx import Presentation  # noqa: E402

from app import pipeline  # noqa: E402
from app.main import app  # noqa: E402
from app.sm2 import sm2  # noqa: E402

pipeline._captions = lambda url: None  # no network in tests
command.upgrade(Config(str(Path(__file__).parents[1] / "alembic.ini")), "head")  # also proves migrations apply
c = TestClient(app)


def auth(email, role):
    r = c.post("/auth/register", json={"email": email, "name": email[:2], "password": "secret1", "role": role})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_sm2():
    assert sm2(5) == (2.6, 1, 1)
    e, i, r = sm2(4, 2.5, 1, 1)
    assert (i, r) == (6, 2)
    assert sm2(4, e, i, r)[1] == round(6 * e)
    assert sm2(1, 2.5, 30, 5)[1:] == (1, 0)
    assert sm2(0, 1.3)[0] == 1.3
    with pytest.raises(ValueError):
        sm2(6)


def test_end_to_end():
    t, st = auth("gv@x.vn", "teacher"), auth("sv@x.vn", "student")
    cls = c.post("/classes", json={"name": "CTDL"}, headers=t).json()
    assert c.post("/classes/join", json={"code": cls["code"].lower()}, headers=st).status_code == 200

    lec = c.post(
        f"/classes/{cls['id']}/lectures", data={"title": "BST", "youtube_url": "https://youtu.be/x"}, headers=t
    ).json()
    assert lec["status"] == "ready" and lec["progress"] == 100
    assert lec["mock"] is True  # mock speech + mock LLM are flagged

    # unpublished lectures are hidden from students
    assert c.get(f"/lectures/{lec['id']}", headers=st).status_code == 403
    cards = c.get(f"/lectures/{lec['id']}/flashcards", headers=t).json()
    assert cards and c.delete(f"/flashcards/{cards[0]['id']}", headers=t).status_code == 204
    c.post(f"/lectures/{lec['id']}/publish", headers=t)

    due = c.get("/review/due", headers=st).json()
    assert len(due) == len(cards) - 1
    assert c.post("/review", json={"flashcard_id": due[0]["id"], "quality": 5}, headers=st).json()["interval"] == 1
    assert len(c.get("/review/due", headers=st).json()) == len(due) - 1

    quiz = c.get(f"/lectures/{lec['id']}/quiz", headers=st).json()
    assert "answer_index" not in quiz[0]
    key = next(
        q["answer_index"] for q in c.get(f"/lectures/{lec['id']}/quiz", headers=t).json() if q["id"] == quiz[0]["id"]
    )
    wrong = (key + 1) % len(quiz[0]["options"])
    assert c.post(f"/quiz/{quiz[0]['id']}/answer", json={"choice": wrong}, headers=st).json()["correct"] is False

    stats = c.get(f"/classes/{cls['id']}/stats", headers=t).json()
    assert stats["questions"][0]["wrong_rate"] == 1
    assert stats["students"][0]["reviews"] == 1
    assert c.get(f"/classes/{cls['id']}/stats", headers=st).status_code == 403

    csv = c.get(f"/lectures/{lec['id']}/export.csv", headers=st).text
    assert "#separator:semicolon" in csv
    apkg = c.get(f"/lectures/{lec['id']}/export.apkg", headers=st).content
    assert "collection.anki2" in zipfile.ZipFile(io.BytesIO(apkg)).namelist()


def test_slide_upload_uses_extracted_text():
    t = auth("gv2@x.vn", "teacher")
    cls = c.post("/classes", json={"name": "OOP"}, headers=t).json()
    deck = Presentation()
    for text in ["Tính đóng gói", "", "Tính kế thừa"]:  # blank slide is skipped
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = text
    buf = io.BytesIO()
    deck.save(buf)
    lec = c.post(
        f"/classes/{cls['id']}/lectures",
        data={"title": "Slide"},
        files={"file": ("bai1.pptx", buf.getvalue())},
        headers=t,
    ).json()
    assert lec["source_type"] == "document" and lec["status"] == "ready", lec
    segs = c.get(f"/lectures/{lec['id']}/transcript", headers=t).json()
    assert [(s["start"], s["text"]) for s in segs] == [(1, "Tính đóng gói"), (3, "Tính kế thừa")]


def test_generate_quiz_with_options():
    t = auth("gv3@x.vn", "teacher")
    cls = c.post("/classes", json={"name": "Q"}, headers=t).json()
    lec = c.post(
        f"/classes/{cls['id']}/lectures", data={"title": "L", "youtube_url": "https://youtu.be/y"}, headers=t
    ).json()
    body = {"count": 5, "difficulty": "hard", "kinds": ["true_false", "calculation"], "replace": True}
    quiz = c.post(f"/lectures/{lec['id']}/quiz/generate", json=body, headers=t).json()
    assert len(quiz) == 5
    assert {q["kind"] for q in quiz} == {"true_false", "calculation"}
    assert all(q["difficulty"] == "hard" for q in quiz)
    assert all(q["options"] == ["Đúng", "Sai"] for q in quiz if q["kind"] == "true_false")
    bad = c.post(f"/lectures/{lec['id']}/quiz/generate", json={**body, "count": 99}, headers=t)
    assert bad.status_code == 422


def test_upload_with_quiz_options():
    t = auth("gv4@x.vn", "teacher")
    cls = c.post("/classes", json={"name": "U"}, headers=t).json()
    form = {"title": "L", "youtube_url": "https://youtu.be/z", "quiz_count": "4", "quiz_difficulty": "easy"}
    lec = c.post(f"/classes/{cls['id']}/lectures", data={**form, "quiz_kinds": "true_false"}, headers=t).json()
    quiz = c.get(f"/lectures/{lec['id']}/quiz", headers=t).json()
    assert len(quiz) == 4 and {(q["kind"], q["difficulty"]) for q in quiz} == {("true_false", "easy")}
    bad = c.post(f"/classes/{cls['id']}/lectures", data={**form, "quiz_kinds": "essay"}, headers=t)
    assert bad.status_code == 422
