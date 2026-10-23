"""Text extraction for slide/PDF lectures. A document segment's `start` is its 1-based page/slide number."""

from pathlib import Path

from ai.speech import Segment

DOC_SUFFIXES = {".pdf", ".pptx"}


def extract_pages(path: Path) -> list[Segment]:
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader  # noqa: PLC0415

        pages = [p.extract_text() or "" for p in PdfReader(path).pages]
    else:
        from pptx import Presentation  # noqa: PLC0415

        pages = [
            "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            for slide in Presentation(path).slides
        ]
    segs = [Segment(i, i, " ".join(t.split())) for i, t in enumerate(pages, 1) if t.strip()]
    if not segs:
        raise ValueError("Không trích được chữ từ tài liệu (có thể là bản scan, cần OCR)")
    return segs
