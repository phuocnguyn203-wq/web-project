import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from ai.service import app as service
from ai.service import pipeline

TOKEN = {"X-Internal-Token": "dev-internal-token"}
client = TestClient(service.app)


@pytest.fixture
def posted(monkeypatch):
    """Capture what the worker reports to the Express API instead of sending it."""
    calls = []
    monkeypatch.setattr(pipeline, "_post", lambda path, body: calls.append((path, body)))
    monkeypatch.setattr(pipeline, "_captions", lambda url: None)  # no network in tests
    return calls


def test_pipeline_reports_progress_then_result(posted):
    quiz = {"count": 4, "difficulty": "easy", "kinds": ["true_false"]}
    pipeline.process_lecture(
        {"lecture_id": 7, "source_type": "youtube", "source_url": "https://youtu.be/x", "quiz": quiz}
    )
    paths = [p for p, _ in posted]
    assert paths[:-1] == ["/internal/lectures/7/status"] * 4 and paths[-1] == "/internal/lectures/7/result"
    assert [b["status"] for _, b in posted[:-1]] == ["transcribing", "summarizing", "generating", "generating"]
    result = posted[-1][1]
    assert result["mock"] is True and result["engine"] == "mock"
    assert result["segments"] and result["summary"] and result["flashcards"]
    assert len(result["quiz"]) == 4 and {(q["kind"], q["difficulty"]) for q in result["quiz"]} == {
        ("true_false", "easy")
    }


def test_slides_use_extracted_text(posted, tmp_path):
    deck = Presentation()
    for text in ["Tính đóng gói", "", "Tính kế thừa"]:  # blank slide is skipped
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = text
    deck.save(tmp_path / "bai1.pptx")
    pipeline.process_lecture({"lecture_id": 1, "source_type": "document", "file_path": str(tmp_path / "bai1.pptx")})
    result = posted[-1][1]
    assert result["engine"] == "document"
    assert [(s["start"], s["text"]) for s in result["segments"]] == [(1, "Tính đóng gói"), (3, "Tính kế thừa")]


def test_failure_is_reported(posted):
    with pytest.raises(FileNotFoundError):
        pipeline.process_lecture({"lecture_id": 2, "source_type": "document", "file_path": "/nope.pdf"})
    assert posted[-1] == ("/internal/lectures/2/status", {"status": "failed", "error": posted[-1][1]["error"]})


def test_internal_token_required():
    assert client.get("/health").status_code == 401
    assert client.get("/health", headers=TOKEN).json() == {"ok": True}


def test_quiz_endpoint_honours_options():
    segs = [{"start": 0, "end": 5, "text": "cây nhị phân tìm kiếm có độ phức tạp O log n"}]
    body = {"segments": segs, "count": 3, "difficulty": "hard", "kinds": ["theory", "true_false"]}
    quiz = client.post("/quiz", json=body, headers=TOKEN).json()
    assert len(quiz) == 3 and all(q["difficulty"] == "hard" for q in quiz)
    assert client.post("/quiz", json={**body, "count": 99}, headers=TOKEN).status_code == 422


def test_jobs_run_in_process_with_sync_jobs(monkeypatch):
    monkeypatch.setenv("SYNC_JOBS", "1")
    ran = []
    monkeypatch.setattr(service, "process_lecture", ran.append)
    job = {"lecture_id": 3, "source_type": "youtube", "source_url": "https://youtu.be/x"}
    assert client.post("/jobs", json=job, headers=TOKEN).status_code == 202
    assert ran[0]["lecture_id"] == 3 and ran[0]["quiz"] is None


def test_apkg_export_is_an_anki_package():
    cards = [{"id": 1, "front": "BST là gì?", "back": "Cây nhị phân tìm kiếm", "source_start": 75}]
    body = {"lecture_id": 1, "title": "BST", "source_type": "youtube", "cards": cards}
    r = client.post("/export/apkg", json=body, headers=TOKEN)
    assert r.status_code == 200
    assert "collection.anki2" in zipfile.ZipFile(io.BytesIO(r.content)).namelist()
